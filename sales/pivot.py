"""產品 × 日期 交叉統計表。

版面（完全照需求）：
  A 欄 = 產品名稱（每列一個產品）
  B 欄起 = 日期（每個出現過的日期一欄，由小到大）
  交叉儲存格 = 該產品在該日期的銷售總量
  最後一欄「合計」= 該產品的總銷量；最後一列「合計」= 當日所有產品總量。
"""

from __future__ import annotations

import io
from collections import defaultdict
from typing import Optional


def build_pivot(records: list[dict]) -> dict:
    """回傳 {dates, products, matrix, row_totals, col_totals, grand_total}。

    products 為 [{name, code}]；matrix[product_name][date] = 數量總和。
    """
    dates: set[str] = set()
    prod_meta: dict[str, str] = {}          # name -> code
    matrix: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))

    for r in records:
        name = r.get("product_name") or "(未指定)"
        d = r.get("date") or ""
        qty = _num(r.get("quantity"))
        dates.add(d)
        prod_meta.setdefault(name, r.get("product_code", "") or "")
        if r.get("product_code"):
            prod_meta[name] = r.get("product_code")
        matrix[name][d] += qty

    sorted_dates = sorted(d for d in dates if d)
    products = sorted(prod_meta.keys())

    row_totals = {p: sum(matrix[p].values()) for p in products}
    col_totals = {d: sum(matrix[p].get(d, 0) for p in products) for d in sorted_dates}
    grand_total = sum(row_totals.values())

    return {
        "dates": sorted_dates,
        "products": [{"name": p, "code": prod_meta.get(p, "")} for p in products],
        "matrix": {p: {d: matrix[p].get(d, 0) for d in sorted_dates} for p in products},
        "row_totals": row_totals,
        "col_totals": col_totals,
        "grand_total": grand_total,
    }


def _num(v) -> float:
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return 0


def to_xlsx(records: list[dict]) -> bytes:
    """產生 Excel（.xlsx）位元組內容。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    piv = build_pivot(records)
    dates = piv["dates"]
    products = piv["products"]

    wb = Workbook()
    ws = wb.active
    ws.title = "銷售統計"

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    total_fill = PatternFill("solid", fgColor="FCE4D6")
    bold = Font(bold=True)
    center = Alignment(horizontal="center")

    # 表頭：A=產品名稱, B..=日期, 末欄=合計
    header = ["產品名稱"] + dates + ["合計"]
    ws.append(header)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=1, column=col)
        c.font = bold
        c.fill = header_fill
        c.alignment = center

    # 內容列
    for p in products:
        label = f"{p['name']}（{p['code']}）" if p["code"] else p["name"]
        row = [label]
        for d in dates:
            row.append(piv["matrix"][p["name"]].get(d, 0) or "")
        row.append(piv["row_totals"][p["name"]])
        ws.append(row)
        ws.cell(row=ws.max_row, column=len(header)).font = bold

    # 合計列
    total_row = ["合計"] + [piv["col_totals"].get(d, 0) for d in dates] + [piv["grand_total"]]
    ws.append(total_row)
    for col in range(1, len(header) + 1):
        c = ws.cell(row=ws.max_row, column=col)
        c.font = bold
        c.fill = total_fill

    # 欄寬與凍結
    ws.column_dimensions["A"].width = 22
    for i in range(2, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 12
    ws.freeze_panes = "B2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(records: list[dict]) -> str:
    import csv as _csv
    import io as _io

    piv = build_pivot(records)
    dates = piv["dates"]
    buf = _io.StringIO()
    buf.write("﻿")  # BOM，讓 Excel 正確辨識 UTF-8
    w = _csv.writer(buf)
    w.writerow(["產品名稱"] + dates + ["合計"])
    for p in piv["products"]:
        label = f"{p['name']}（{p['code']}）" if p["code"] else p["name"]
        w.writerow([label] + [piv["matrix"][p["name"]].get(d, 0) for d in dates] + [piv["row_totals"][p["name"]]])
    w.writerow(["合計"] + [piv["col_totals"].get(d, 0) for d in dates] + [piv["grand_total"]])
    return buf.getvalue()
