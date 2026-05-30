"""盤點差異報表 — 匯出 Excel / CSV。

內容：批次資訊 + 差異明細（帳面 / 實盤 / 差異 / 原因 / 狀態 / 盤點人員），
可依「類別」或「差異大小（絕對值）」排序，並附結案後的調整紀錄。
"""

from __future__ import annotations

import io

# 報表欄位（明細）
COLUMNS = ["商品編號", "商品名稱", "類別", "規格", "帳面庫存", "實盤庫存",
           "差異數量", "差異原因", "狀態", "盤點人員"]


def _num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def _abs_diff(it) -> float:
    d = it.get("diff_qty")
    try:
        return abs(float(d))
    except (TypeError, ValueError):
        return -1.0  # 未盤（diff 為 None）排最後


def select_rows(items: list[dict], only_diff: bool = True, sort: str = "category") -> list[dict]:
    """過濾與排序差異品項。

    only_diff=True 只取差異/已覆核且差異≠0 的品項；False 取全部已盤。
    sort='category' 依 類別→編號；sort='diff' 依差異絕對值由大到小。
    """
    rows = []
    for it in items:
        if it.get("actual_qty") is None:
            continue  # 未盤不列入
        diff = it.get("diff_qty") or 0
        if only_diff and diff == 0:
            continue
        rows.append(it)
    if sort == "diff":
        rows.sort(key=lambda it: (-_abs_diff(it), it.get("product_code") or ""))
    else:
        rows.sort(key=lambda it: (it.get("category_snapshot") or "", it.get("product_code") or ""))
    return rows


def summarize(rows: list[dict]) -> dict:
    pos = sum(1 for it in rows if (it.get("diff_qty") or 0) > 0)
    neg = sum(1 for it in rows if (it.get("diff_qty") or 0) < 0)
    abs_sum = sum(_abs_diff(it) for it in rows if _abs_diff(it) >= 0)
    return {"count": len(rows), "over": pos, "short": neg, "abs_sum": _num(abs_sum)}


def _row_values(it) -> list:
    return [
        it.get("product_code", ""),
        it.get("name_snapshot", ""),
        it.get("category_snapshot", ""),
        it.get("spec_snapshot", ""),
        _num(it.get("book_qty_snapshot")),
        _num(it.get("actual_qty")),
        _num(it.get("diff_qty")),
        it.get("diff_reason", ""),
        it.get("status", ""),
        it.get("counter", ""),
    ]


def to_xlsx(batch: dict, items: list[dict], adjustments: list[dict] | None = None,
            only_diff: bool = True, sort: str = "category") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = select_rows(items, only_diff=only_diff, sort=sort)
    summary = summarize(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "差異報表"

    bold = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    over_fill = PatternFill("solid", fgColor="FCE4D6")    # 盤盈（淺橘）
    short_fill = PatternFill("solid", fgColor="FCE7E7")   # 盤虧（淺紅）
    center = Alignment(horizontal="center")

    # 標題與摘要
    title = f"盤點差異報表 — {batch.get('title') or batch.get('id', '')}"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"盤點日期 {batch.get('count_date', '')} ·"
               f" 狀態 {batch.get('status', '')} ·"
               f" 差異 {summary['count']} 項（盤盈 {summary['over']}、盤虧 {summary['short']}）·"
               f" 差異量合計 {summary['abs_sum']}"])
    ws.append([])

    head_row = ws.max_row + 1
    ws.append(COLUMNS)
    for col in range(1, len(COLUMNS) + 1):
        c = ws.cell(row=head_row, column=col)
        c.font = bold; c.fill = header_fill; c.alignment = center

    for it in rows:
        ws.append(_row_values(it))
        diff = it.get("diff_qty") or 0
        if diff:
            fill = over_fill if diff > 0 else short_fill
            ws.cell(row=ws.max_row, column=7).fill = fill
            ws.cell(row=ws.max_row, column=7).font = bold

    # 調整紀錄（結案後）
    if adjustments:
        ws.append([])
        ws.append(["調整紀錄（結案後）"]); ws.cell(row=ws.max_row, column=1).font = bold
        adj_head = ws.max_row + 1
        adj_cols = ["時間", "商品編號", "原實盤", "調整後", "原因", "調整人"]
        ws.append(adj_cols)
        for col in range(1, len(adj_cols) + 1):
            ws.cell(row=adj_head, column=col).font = bold
            ws.cell(row=adj_head, column=col).fill = header_fill
        for a in adjustments:
            ws.append([a.get("created_at", ""), a.get("product_code", ""),
                       _num(a.get("old_actual")), _num(a.get("new_actual")),
                       a.get("reason", ""), a.get("created_by", "")])

    widths = [16, 22, 10, 14, 10, 10, 10, 18, 8, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{head_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(batch: dict, items: list[dict], adjustments: list[dict] | None = None,
           only_diff: bool = True, sort: str = "category") -> str:
    import csv as _csv

    rows = select_rows(items, only_diff=only_diff, sort=sort)
    summary = summarize(rows)
    buf = io.StringIO()
    buf.write("﻿")  # BOM，讓 Excel 正確辨識 UTF-8
    w = _csv.writer(buf)
    w.writerow([f"盤點差異報表 — {batch.get('title') or batch.get('id', '')}"])
    w.writerow([f"盤點日期 {batch.get('count_date', '')}", f"狀態 {batch.get('status', '')}",
                f"差異 {summary['count']} 項", f"盤盈 {summary['over']}", f"盤虧 {summary['short']}",
                f"差異量合計 {summary['abs_sum']}"])
    w.writerow([])
    w.writerow(COLUMNS)
    for it in rows:
        w.writerow(_row_values(it))
    if adjustments:
        w.writerow([])
        w.writerow(["調整紀錄（結案後）"])
        w.writerow(["時間", "商品編號", "原實盤", "調整後", "原因", "調整人"])
        for a in adjustments:
            w.writerow([a.get("created_at", ""), a.get("product_code", ""),
                        _num(a.get("old_actual")), _num(a.get("new_actual")),
                        a.get("reason", ""), a.get("created_by", "")])
    return buf.getvalue()
